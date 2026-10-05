# -*- coding: utf-8 -*-
"""회사 견적서 템플릿(assets/quote_template.xlsx)을 채워 새 견적서를 만든다.

템플릿에 로고/직인 이미지가 들어 있어 openpyxl로 열고 저장하면 이미지가 유실된다.
그래서 zip 안의 sheet XML에서 지정 셀만 교체하는 방식을 쓴다. 수식은 건드리지 않고
캐시값만 갱신하므로 재계산이 필요 없다.

사용법:
  python fill_quote.py --out "경로/견적서_대상.xlsx" --date 2026-01-15 \
      --number "견적2026-J001" --client "㈜테스트" --title "OO 시스템 v1.0" \
      --amount 5000000 --lines-file lines.txt [--travel 0] [--overwrite]

lines.txt: 견적내용 칸(C17~C26, 최대 10줄)에 들어갈 줄들. UTF-8, 한 줄에 한 항목.
보통 이런 형식을 쓴다(기존 견적서 관례):
  - 업   체   명 : ㈜테스트
  - 구         분 : 일반 시험성적서
  - 시험 항목 1 : ...
amount는 공급가액(VAT 별도) 정수. --travel(출장비)을 주면 합계에 더해진다.
"""
import argparse
import datetime
import os
import re
import sys
import zipfile
from xml.sax.saxutils import escape

TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "quote_template.xlsx")
LINE_CELLS = [f"C{r}" for r in range(17, 27)]  # 견적내용 줄들


def replace_cell(xml, ref, new_inner, keep_formula=False):
    """<c r="ref" ...> 요소를 찾아 내용을 교체한다. 스타일(s=)은 보존."""
    # [^>]*? 는 게으르게: 탐욕적으로 두면 자기닫힘 셀의 '/'까지 삼켜
    # 다음 셀들까지 매칭이 번져 템플릿을 망가뜨린다.
    pat = rf'<c r="{ref}"( s="\d+")?[^>]*?(?:/>|>.*?</c>)'
    m = re.search(pat, xml)
    if not m:
        raise SystemExit(f"템플릿에서 셀 {ref}를 찾지 못했습니다. 템플릿이 바뀌었는지 확인하세요.")
    style = m.group(1) or ""
    return xml[: m.start()] + f'<c r="{ref}"{style}{new_inner}' + xml[m.end():], None


def set_text(xml, ref, text):
    inner = f' t="inlineStr"><is><t xml:space="preserve">{escape(text)}</t></is></c>'
    xml, _ = replace_cell(xml, ref, inner)
    return xml


def set_number(xml, ref, value):
    xml, _ = replace_cell(xml, ref, f"><v>{value}</v></c>")
    return xml


def set_empty(xml, ref):
    xml, _ = replace_cell(xml, ref, "/>")
    return xml


def set_formula_cache(xml, ref, value):
    """수식 셀의 캐시값(<v>)만 갱신한다. 수식 자체는 보존."""
    pat = rf'(<c r="{ref}"[^>]*><f[^>]*(?:/>|>.*?</f>))(?:<v>[^<]*</v>)?(</c>)'
    new, n = re.subn(pat, rf"\g<1><v>{value}</v>\g<2>", xml)
    if n != 1:
        raise SystemExit(f"수식 셀 {ref} 캐시 갱신 실패 (매칭 {n}건)")
    return new


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", default=TEMPLATE)
    ap.add_argument("--out", required=True)
    ap.add_argument("--date", required=True, help="YYYY-MM-DD")
    ap.add_argument("--number", required=True, help="견적번호 예: 견적2026-J001")
    ap.add_argument("--client", required=True, help="수신(대상 기업명)")
    ap.add_argument("--title", required=True, help="견적명")
    ap.add_argument("--amount", required=True, type=int, help="공급가액(VAT별도)")
    ap.add_argument("--lines-file", required=True, help="견적내용 줄들(UTF-8, 줄당 1항목, 최대 10줄)")
    ap.add_argument("--travel", type=int, default=0, help="출장비(공급가액에 합산)")
    ap.add_argument("--overwrite", action="store_true")
    a = ap.parse_args()

    if os.path.exists(a.out) and not a.overwrite:
        raise SystemExit(f"이미 존재합니다: {a.out} (덮어쓰려면 --overwrite)")
    with open(a.lines_file, encoding="utf-8-sig") as f:
        lines = [ln.rstrip("\n") for ln in f if ln.strip()]
    if len(lines) > len(LINE_CELLS):
        raise SystemExit(f"견적내용은 최대 {len(LINE_CELLS)}줄까지입니다 (현재 {len(lines)}줄)")

    d = datetime.date.fromisoformat(a.date)
    total = a.amount + a.travel

    zin = zipfile.ZipFile(a.template)
    xml = zin.read("xl/worksheets/sheet1.xml").decode("utf-8")

    xml = set_text(xml, "A6", f"수     신 : {a.client}")
    xml = set_text(xml, "A8", f"견 적 일 : {d.year}년 {d.month:02d}월 {d.day:02d}일")
    xml = set_text(xml, "I9", f"[견적번호 : {a.number}]")
    xml = set_text(xml, "A13", f"견 적 명 : {a.title}")
    for i, ref in enumerate(LINE_CELLS):
        xml = set_text(xml, ref, lines[i]) if i < len(lines) else set_empty(xml, ref)
    xml = set_number(xml, "I30", a.amount)          # 시험성적서 발급 금액
    if a.travel:
        xml = set_number(xml, "E31", a.travel)      # 출장비 입력칸
    xml = set_formula_cache(xml, "I31", a.travel)   # 출장비(=E31)
    xml = set_formula_cache(xml, "I32", total)      # 산출내역 합계
    xml = set_formula_cache(xml, "E9", total)       # 상단 견적금액(=I32)
    xml = set_formula_cache(xml, "G41", total)      # 최종합계(=I32)

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    with zipfile.ZipFile(a.out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    zin.close()

    # 검증: 값이 제대로 들어갔고 이미지가 보존되었는지
    import openpyxl
    wb = openpyxl.load_workbook(a.out, data_only=True)
    ws = wb["견적서"]
    ok = (a.client in str(ws["A6"].value) and a.number in str(ws["I9"].value)
          and ws["E9"].value == total)
    media = [n for n in zipfile.ZipFile(a.out).namelist() if n.startswith("xl/media/")]
    if not ok or not media:
        raise SystemExit(f"검증 실패: values_ok={ok}, media={len(media)}")
    print(f"OK: {a.out} 생성 (공급가액 {total:,}원, VAT포함 {int(total*1.1):,}원, 이미지 {len(media)}개 보존)")


if __name__ == "__main__":
    sys.exit(main())
