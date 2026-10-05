# -*- coding: utf-8 -*-
"""견적서발행대장.xlsx에 견적 1건을 기입하고, 매출소득관리.xlsx를 동기화한다.

왜 이렇게 복잡한가:
- 매출소득관리.xlsx는 발행대장을 `=[1]견적서발행!A3` 형태의 외부 링크로 참조한다.
  openpyxl로 열어서 저장하면 외부 링크 캐시가 사라져 파일이 망가지므로,
  zip 안의 XML만 직접 수정한다(외부 링크 캐시 갱신 + 수식 캐시 제거).
- 수식 캐시를 제거한 뒤 LibreOffice headless 변환으로 재계산한다.
  (캐시가 남아 있으면 LibreOffice가 재계산을 건너뛴다)
- LibreOffice가 링크 경로를 ../ 상대경로로 바꿔놓으므로, 마지막에 같은 폴더
  기준 상대경로(파일명만)로 복원한다. 사용자는 두 파일을 항상 같은 폴더에 둔다.

사용법:
  기입+동기화: python register_quote.py --base DIR --date 2026-01-15 \
                 --number "견적2026-J001" --client "㈜테스트" \
                 --title "OO 시스템 v1.0" --amount 5000000 [--note "..."]
  동기화만:    python register_quote.py --base DIR --sync-only
                 (발행대장을 수동으로 고친 뒤 매출소득관리만 다시 맞출 때)

--base: 견적서발행대장.xlsx 와 매출소득관리.xlsx 가 있는 폴더.
--amount: 공급가액(VAT별도). 부가세/합계는 대장의 기존 수식이 계산한다.
"""
import argparse
import datetime
import glob
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from xml.sax.saxutils import escape

LEDGER = "견적서발행대장.xlsx"
INCOME = "매출소득관리.xlsx"
LEDGER_SHEET = "견적서발행"


def find_soffice():
    cand = [shutil.which("soffice"),
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]
    for c in cand:
        if c and os.path.exists(c):
            return c
    raise SystemExit("LibreOffice(soffice)를 찾을 수 없습니다. 재계산에 필요합니다.")


def lo_recalc(path):
    """LibreOffice로 파일을 같은 형식으로 변환해 모든 수식을 재계산하고 제자리 교체한다.
    주의: 수식 셀에 캐시값이 남아 있으면 재계산되지 않는다. 호출 전에 캐시를 제거할 것."""
    soffice = find_soffice()
    with tempfile.TemporaryDirectory() as td:
        r = subprocess.run(
            [soffice, "--headless", "--norestore",
             "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", td, path],
            capture_output=True, timeout=120)
        out = os.path.join(td, os.path.basename(path))
        if r.returncode != 0 or not os.path.exists(out):
            raise SystemExit(f"LibreOffice 변환 실패: {r.stderr.decode(errors='replace')[:300]}")
        shutil.copy(out, path)


def strip_formula_caches(sheet_xml):
    """수식 셀의 캐시값 <v>를 제거해 강제 재계산되게 한다."""
    return re.sub(r"(<f[^>]*(?:/>|>.*?</f>))<v>[^<]*</v>", r"\1", sheet_xml, flags=re.S)


def check_errors(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    errs = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#") and c.value.rstrip("!?") in (
                        "#REF", "#NAME", "#VALUE", "#DIV/0", "#N/A", "#NULL", "#NUM"):
                    errs.append(f"{ws.title}!{c.coordinate}={c.value}")
    return errs


def append_to_ledger(base, a):
    import openpyxl
    path = os.path.join(base, LEDGER)
    wb = openpyxl.load_workbook(path)
    ws = wb[LEDGER_SHEET]
    # 중복 견적번호 방지
    for r in range(3, 33):
        if ws.cell(row=r, column=3).value == a.number:
            raise SystemExit(f"견적번호 {a.number}는 이미 {r}행에 있습니다.")
    # 빈 행 찾기: B(발행일)~E(견적명)가 모두 비어 있는 첫 행. F열은 0이 미리 채워져 있을 수 있으니 무시.
    target = None
    for r in range(3, 33):
        if all(ws.cell(row=r, column=c).value in (None, "") for c in (2, 3, 4, 5)):
            target = r
            break
    if target is None:
        raise SystemExit("발행대장(3~32행)에 빈 행이 없습니다. 행을 늘린 뒤 다시 실행하세요.")
    d = datetime.date.fromisoformat(a.date)
    cell = ws.cell(row=target, column=2, value=d)
    cell.number_format = "yyyy-mm-dd"
    if a.number:
        ws.cell(row=target, column=3, value=a.number)
    ws.cell(row=target, column=4, value=a.client)
    ws.cell(row=target, column=5, value=a.title)
    ws.cell(row=target, column=6, value=a.amount)
    if a.note:
        ws.cell(row=target, column=9, value=a.note)
    wb.save(path)          # openpyxl 저장으로 수식 캐시가 비워짐
    lo_recalc(path)        # → LibreOffice가 전부 재계산
    errs = check_errors(path)
    if errs:
        raise SystemExit(f"발행대장 수식 오류: {errs[:10]}")
    return target


def build_link_cache(base):
    """발행대장의 현재 값(A~I, 3~34행)으로 외부 링크 캐시 XML을 만든다."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.datetime import to_excel
    wb = openpyxl.load_workbook(os.path.join(base, LEDGER), data_only=True)
    ws = wb[LEDGER_SHEET]
    rows = []
    for r in range(3, 35):
        cells = []
        for c in range(1, 10):
            v = ws.cell(row=r, column=c).value
            if v is None or v == "":
                continue
            ref = f"{get_column_letter(c)}{r}"
            if isinstance(v, str):
                cells.append(f'<cell r="{ref}" t="str"><v>{escape(v)}</v></cell>')
            else:
                if isinstance(v, (datetime.datetime, datetime.date)):
                    v = to_excel(v)
                if isinstance(v, float) and v == int(v):
                    v = int(v)
                cells.append(f'<cell r="{ref}"><v>{v!r}</v></cell>')
        if cells:
            rows.append(f'<row r="{r}">{"".join(cells)}</row>')
    return "".join(rows)


def sync_income(base):
    """매출소득관리.xlsx의 외부 링크 캐시를 갱신하고 재계산한다. zip 수준으로만 수정."""
    path = os.path.join(base, INCOME)
    cache_rows = build_link_cache(base)
    zin = zipfile.ZipFile(path)
    names = zin.namelist()

    link_files = [n for n in names if re.fullmatch(r"xl/externalLinks/externalLink\d+\.xml", n)]
    replacements = {}
    for lf in link_files:
        xml = zin.read(lf).decode("utf-8")
        m = re.search(r"<sheetNames>(.*?)</sheetNames>", xml, re.S)
        if not m or f'val="{LEDGER_SHEET}"' not in m.group(1):
            continue
        sheet_idx = [re.search(r'val="([^"]*)"', s).group(1)
                     for s in re.findall(r"<sheetName [^>]*/>", m.group(1))].index(LEDGER_SHEET)
        new_sd = f'<sheetData sheetId="{sheet_idx}">{cache_rows}</sheetData>'
        xml, n = re.subn(rf'<sheetData sheetId="{sheet_idx}"(?:/>|>.*?</sheetData>)', new_sd, xml,
                         flags=re.S)
        if n != 1:
            raise SystemExit(f"{lf}: 캐시 교체 실패")
        replacements[lf] = xml.encode("utf-8")
    if not replacements:
        raise SystemExit(f"매출소득관리에서 '{LEDGER_SHEET}' 외부 링크를 찾지 못했습니다.")

    # 수식 캐시 제거(강제 재계산) + calcChain 제거(Excel 저장분에 남아 있을 수 있음)
    for n in names:
        if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n):
            replacements[n] = strip_formula_caches(zin.read(n).decode("utf-8")).encode("utf-8")
    drop = set()
    if "xl/calcChain.xml" in names:
        drop.add("xl/calcChain.xml")
        ct = zin.read("[Content_Types].xml").decode("utf-8")
        replacements["[Content_Types].xml"] = re.sub(
            r'<Override PartName="/xl/calcChain\.xml"[^>]*/>', "", ct).encode("utf-8")
        wr = zin.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        replacements["xl/_rels/workbook.xml.rels"] = re.sub(
            r'<Relationship[^>]*Target="calcChain\.xml"[^>]*/>', "", wr).encode("utf-8")

    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename in drop:
                continue
            zout.writestr(item, replacements.get(item.filename, zin.read(item.filename)))
    zin.close()
    shutil.move(tmp, path)

    lo_recalc(path)

    # LibreOffice가 바꾼 링크 경로를 같은 폴더 상대경로(파일명만)로 복원
    zin = zipfile.ZipFile(path)
    rel_fix = {}
    for n in zin.namelist():
        if re.fullmatch(r"xl/externalLinks/_rels/externalLink\d+\.xml\.rels", n):
            rels = zin.read(n).decode("utf-8")
            fixed = re.sub(r'Target="[^"]*견적서발행대장\.xlsx"', f'Target="{LEDGER}"', rels)
            if fixed != rels:
                rel_fix[n] = fixed.encode("utf-8")
    if rel_fix:
        tmp = path + ".tmp"
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, rel_fix.get(item.filename, zin.read(item.filename)))
        zin.close()
        shutil.move(tmp, path)
    else:
        zin.close()

    errs = check_errors(path)
    if errs:
        raise SystemExit(f"매출소득관리 수식 오류: {errs[:10]}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", required=True, help="두 관리 파일이 있는 폴더")
    ap.add_argument("--sync-only", action="store_true", help="매출소득관리 동기화만 수행")
    ap.add_argument("--date")
    ap.add_argument("--number", default="")
    ap.add_argument("--client")
    ap.add_argument("--title")
    ap.add_argument("--amount", type=int)
    ap.add_argument("--note", default="")
    a = ap.parse_args()

    base = os.path.abspath(a.base)
    for f in (LEDGER, INCOME):
        if not os.path.exists(os.path.join(base, f)):
            raise SystemExit(f"{f}가 {base}에 없습니다.")

    if a.sync_only:
        sync_income(base)
        print("OK: 매출소득관리 동기화 완료")
        return
    for req in ("date", "client", "title", "amount"):
        if getattr(a, req) in (None, ""):
            raise SystemExit(f"--{req} 인자가 필요합니다 (또는 --sync-only)")
    row = append_to_ledger(base, a)
    sync_income(base)
    vat = a.amount // 10
    print(f"OK: 발행대장 {row}행 기입(발행순서 {row-2}), 공급가액 {a.amount:,}원 / "
          f"부가세 {vat:,}원 / 합계 {a.amount + vat:,}원. 매출소득관리 동기화 완료.")


if __name__ == "__main__":
    sys.exit(main())
